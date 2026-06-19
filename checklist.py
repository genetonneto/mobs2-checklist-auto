from playwright.sync_api import sync_playwright, Page
import os
import re
from dataclasses import dataclass
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import PatternFill, Font

SESSION_FILE = "session.json"
URL_BASE = "https://system.mobs2.com/"
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/122.0.0.0 Safari/537.36"
)
SCRIPT_ANTI_WEBDRIVER = """
    Object.defineProperty(navigator, 'webdriver', { get: () => undefined });
"""

LIMITE_ATRASO = timedelta(hours=1, minutes=30)
RELATORIO_EXCEL = "relatorio_atraso.xlsx"

COLUNAS_RELATORIO = ["id_unit", "description", "chip", "ultimaposicao"]


# ---------------------------------------------------------------------------
# Sessão / login
# ---------------------------------------------------------------------------

def _novo_contexto(navegador, storage_state=None):
    contexto = navegador.new_context(
        user_agent=USER_AGENT,
        storage_state=storage_state,
    )
    contexto.add_init_script(SCRIPT_ANTI_WEBDRIVER)
    return contexto


def iniciar_sessao(navegador):
    if os.path.exists(SESSION_FILE):
        contexto = _novo_contexto(navegador, storage_state=SESSION_FILE)
        page = contexto.new_page()
        page.goto(URL_BASE)
    else:
        contexto = _novo_contexto(navegador)
        page = contexto.new_page()
        page.goto(URL_BASE)

        page.get_by_role("textbox", name="Email").fill("SEU_EMAIL")
        page.get_by_role("textbox", name="Senha Senha Atual Nova Senha").fill("SUA_SENHA")
        page.get_by_role("button", name="Entrar").click()
        page.wait_for_load_state("networkidle")
        contexto.storage_state(path=SESSION_FILE)

    return page




# ---------------------------------------------------------------------------
# Clientes
# ---------------------------------------------------------------------------

@dataclass
class Cliente:
    nome: str
    busca: str = None

    def __post_init__(self):
        if self.busca is None:
            self.busca = self.nome


# Para inibir um cliente, basta comentar a linha correspondente
CLIENTES = [
    Cliente("ASA BRANCA"),
    # Cliente("BR7"),
    # Cliente("ITAMARACA",         "ITA"),
    # Cliente("RAYMUNDO DA FONTE", "RAYMUNDO"),
    # Cliente("ATALAIA"),
    # Cliente("MB LIMPEZA"),
    # Cliente("BORBOREMA",         "BORBO"),
    # Cliente("M DIAS BRANCO",     "M DIAS"),
    # Cliente("GRUPO SANTA ZITA"),
    # Cliente("DELLYS"),
    # Cliente("SAMBAIBA",          "SAMBAIBA TRANSPORTES"),
    # Cliente("MONTE ALEGRE",      "AGENCIA DE TURISMO MONTE ALEGRE"),
    # Cliente("VALINHOS",          "RAPIDO LUXO - VALINHOS"),
    # Cliente("RAPIDO LUXO CAMPINAS"),
    # Cliente("RS - PAULINIA",     "RAPIDO SUMARE LTDA - PAULINIA"),
    # Cliente("RS - PIRACICABA",   "RAPIDO SUMARE LTDA - PIRACICABA"),
    # Cliente("VB TRANSPORTE"),
    # Cliente("CAMPO LIMPO"),
]


# ---------------------------------------------------------------------------
# Funções de ação
# ---------------------------------------------------------------------------

def selecionar_cliente(page: Page, busca: str) -> None:
    campo = page.get_by_role("combobox").first
    campo.wait_for()
    campo.click()
    campo.type(busca, delay=100)
    page.keyboard.press("Enter")


def verificar_telemetria(page: Page) -> tuple:
    capturas = []

    def on_response(response):
        if "json" in response.headers.get("content-type", ""):
            try:
                body = response.json()
                capturas.append({"url": response.url, "body": body})
            except Exception:
                pass

    page.on("response", on_response)

    page.get_by_text("Cockpit").wait_for(state="visible")
    page.get_by_text("Cockpit").click()
    page.get_by_role("link", name=" Telemetria ").click()

    page.evaluate("document.documentElement.style.zoom = '67%'")
    page.wait_for_timeout(9999)
    page.get_by_label("*").click()

    for _ in range(4):
        page.keyboard.press("ArrowDown")
    page.keyboard.press("Enter")
    page.keyboard.press("Escape")

    gridcell_disponivel = page.get_by_role("gridcell", name=re.compile(r"dispon[ií]vel", re.IGNORECASE)).first
    gridcell_disponivel.click()
    gridcell_disponivel.click()

    page.get_by_role("link", name="Colunas").click()
    page.get_by_role("link", name="Chip Número").click()
    page.mouse.click(page.viewport_size["width"] / 2, page.viewport_size["height"] / 2)

    page.wait_for_timeout(1000)
    for _ in range(200):
        page.mouse.wheel(0, 100)
        page.wait_for_timeout(150)

    page.wait_for_timeout(2000)
    page.remove_listener("response", on_response)

    return processar_capturas_rede(capturas)


def processar_capturas_rede(capturas: list) -> tuple:
    print(f"\n[rede] {len(capturas)} respostas JSON capturadas:")
    for c in capturas:
        tamanho = len(c["body"]) if isinstance(c["body"], list) else "dict"
        print(f"  {tamanho:>6}  {c['url']}")

    melhor_lista = []
    for captura in capturas:
        body = captura["body"]
        candidatos = [body] if isinstance(body, list) else [v for v in body.values() if isinstance(v, list)]
        for candidato in candidatos:
            if len(candidato) > len(melhor_lista):
                melhor_lista = candidato

    if not melhor_lista:
        print("[rede] nenhuma lista de veículos encontrada.")
        return [], []

    print(f"[rede] usando lista com {len(melhor_lista)} itens.")

    if melhor_lista and isinstance(melhor_lista[0], dict):
        cabecalhos = list(melhor_lista[0].keys())
        linhas = [[str(item.get(k, "")) for k in cabecalhos] for item in melhor_lista]
    else:
        cabecalhos = []
        linhas = [[str(v) for v in item] if isinstance(item, list) else [str(item)] for item in melhor_lista]

    return cabecalhos, linhas


def filtrar_veiculos_atrasados(cabecalhos: list, linhas: list) -> list:
    agora = datetime.now()

    idx_ultimaposicao = next(
        (i for i, h in enumerate(cabecalhos) if _normalizar(h) == "ultimaposicao"),
        None,
    )

    if idx_ultimaposicao is None:
        print("[erro] coluna 'ultimaposicao' não encontrada.")
        return []

    atrasados = []
    for linha in linhas:
        try:
            if idx_ultimaposicao >= len(linha):
                continue
            dt = datetime.strptime(str(linha[idx_ultimaposicao]).strip(), "%d/%m/%Y %H:%M")
            atraso = agora - dt
            if atraso >= LIMITE_ATRASO:
                horas = round(atraso.total_seconds() / 3600, 1)
                atrasados.append((linha, horas))
        except Exception:
            continue

    return atrasados


def _normalizar(texto: str) -> str:
    return re.sub(r"[\s_]", "", texto).lower()


def gerar_relatorio_excel(resultados: dict) -> None:
    wb = openpyxl.Workbook()

    # --- Aba de resumo (tabela principal) ---
    ws_resumo = wb.active
    ws_resumo.title = "Resumo"

    fill_alerta = PatternFill(fill_type="solid", fgColor="FF4444")
    fill_ok     = PatternFill(fill_type="solid", fgColor="4CAF50")
    fill_header = PatternFill(fill_type="solid", fgColor="1E3A5F")
    fonte_header = Font(bold=True, color="FFFFFF")
    fonte_alerta = Font(bold=True, color="FFFFFF")
    fonte_ok     = Font(bold=True, color="FFFFFF")

    ws_resumo.append(["Empresa", "Alerta", "Total Atrasados"])
    for cell in ws_resumo[1]:
        cell.fill = fill_header
        cell.font = fonte_header

    for empresa, (_cab, atrasados) in sorted(resultados.items()):
        total_empresa = len(atrasados)
        if total_empresa > 0:
            alerta_texto = "⚠ ALERTA"
            fill_row = fill_alerta
            fonte_row = fonte_alerta
        else:
            alerta_texto = "✔ OK"
            fill_row = fill_ok
            fonte_row = fonte_ok
        ws_resumo.append([empresa, alerta_texto, total_empresa])
        for cell in ws_resumo[ws_resumo.max_row]:
            cell.fill = fill_row
            cell.font = fonte_row

    for col in ws_resumo.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_resumo.column_dimensions[col[0].column_letter].width = max_len + 4

    # --- Aba de detalhes ---
    ws = wb.create_sheet("Veículos Atrasados")

    cabecalhos_globais = next(
        (cab for cab, _ in resultados.values() if cab), []
    )

    alvos = [_normalizar(c) for c in COLUNAS_RELATORIO]
    indices = [i for i, h in enumerate(cabecalhos_globais) if _normalizar(h) in alvos]
    nomes_saida = [cabecalhos_globais[i] for i in indices]

    ws.append(["Empresa"] + nomes_saida + ["Atraso (h)"])
    for cell in ws[1]:
        cell.fill = fill_header
        cell.font = fonte_header

    for empresa, (_cab, atrasados) in sorted(resultados.items()):
        for linha, horas in atrasados:
            celulas = [linha[i] for i in indices if i < len(linha)]
            ws.append([empresa] + celulas + [horas])
            for cell in ws[ws.max_row]:
                cell.fill = fill_alerta
                cell.font = fonte_alerta

    total = sum(len(atrasados) for _, atrasados in resultados.values())
    wb.save(RELATORIO_EXCEL)

    agora = datetime.now()
    print(f"Relatório gerado: {RELATORIO_EXCEL} — {total} veículo(s) atrasado(s)")
    if total > 0:
        print(f"\n{'='*60}")
        print(f"  ALERTA: {total} veículo(s) com último posicionamento")
        print(f"  >= 1h30m atrás ({agora.strftime('%d/%m/%Y %H:%M')} horário do computador).")
        print(f"{'='*60}\n")


# def verificar_localizacao(page: Page) -> None:
#     page.get_by_text("Operação Telemetria").click()
#     page.get_by_role("link", name=" Localização ").click()
#     page.evaluate("document.documentElement.style.zoom = '67%'")
#     page.wait_for_timeout(5000)

#     page.get_by_role("link", name="Todos ").first.click()
#     page.wait_for_timeout(5000)
#     page.get_by_role("link", name="Todos ").nth(1).click()
#     page.wait_for_timeout(5000)
#     page.get_by_role("link", name="Todos ").nth(2).click()
#     page.wait_for_timeout(5000)

#     page.get_by_role("link", name="Selecionar").click()
#     page.wait_for_timeout(500)


# def verificar_relatorios(page: Page) -> None:
#     page.get_by_text("Operação Telemetria").click()
#     page.get_by_role("link", name=" Relatórios ").click()
#     page.evaluate("document.documentElement.style.zoom = '67%'")
#     page.wait_for_timeout(5000)

#     page.get_by_role("link", name=" Mobs2 - Dashboard Icon Dashboard Geral").click()
#     page.evaluate("document.documentElement.style.zoom = '67%'")
#     page.wait_for_timeout(9000)

#     page.get_by_text("Operação Telemetria").click()
#     page.get_by_role("link", name=" Relatórios ").click()
#     page.evaluate("document.documentElement.style.zoom = '67%'")
#     page.wait_for_timeout(5000)

#     page.get_by_role("link", name=" Mobs2 - Dashboard Icon Dashboard de Mapeamento de Eventos").click()
#     page.evaluate("document.documentElement.style.zoom = '67%'")
#     page.wait_for_timeout(9000)




# ---------------------------------------------------------------------------
# Orquestração
# ---------------------------------------------------------------------------

def processar_cliente(page: Page, cliente: Cliente) -> tuple:
    selecionar_cliente(page, cliente.busca)
    cabecalhos, linhas = verificar_telemetria(page)
    atrasados = filtrar_veiculos_atrasados(cabecalhos, linhas)
    return cabecalhos, atrasados
    # verificar_localizacao(page)
    # verificar_relatorios(page)


def main():
    with sync_playwright() as pw:
        navegador = pw.chromium.launch(
            headless=False,
            args=["--disable-blink-features=AutomationControlled"],
        )

        page = iniciar_sessao(navegador)

        resultados = {}
        for cliente in CLIENTES:
            resultados[cliente.nome] = processar_cliente(page, cliente)

        gerar_relatorio_excel(resultados)
        page.pause()


if __name__ == "__main__":
    main()
